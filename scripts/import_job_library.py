from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
API_ROOT = PROJECT_ROOT / "apps" / "api"
if str(API_ROOT) not in sys.path:
    sys.path.insert(0, str(API_ROOT))

from app.core.database import SessionLocal, init_db  # noqa: E402
from app.models.models import JobLibraryEntry  # noqa: E402


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _records(path: Path, sheet_name: str, header_row: int):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        headers = [_text(cell.value) for cell in sheet[header_row]]
        for row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            row = {header: values[index] for index, header in enumerate(headers) if header and index < len(values)}
            if any(_text(value) for value in values):
                yield row_number, row
    finally:
        workbook.close()


def _identity(company: str, title: str) -> str:
    normalized = re.sub(r"\s+", "", f"{company}||{title}").casefold()
    return hashlib.sha1(normalized.encode("utf-8")).hexdigest()[:24]


def _entry(collection: str, source_sheet: str, row_number: int, row: dict, *, core: bool = False) -> dict:
    company = _text(row.get("公司名称") or row.get("公司"))
    title = _text(row.get("推荐投递岗位") or row.get("招聘岗位") or row.get("岗位"))
    source_key = f"all:{row_number}" if collection == "all" else f"progress:{_identity(company, title)}"
    return {
        "source_key": source_key,
        "collection": collection,
        "source_sheet": source_sheet,
        "source_row": row_number,
        "recommendation_level": "核心推荐" if core else ("总库" if collection == "all" else "可投递"),
        "target_direction": _text(row.get("目标方向") or row.get("主要方向")),
        "other_directions": _text(row.get("其他匹配方向")),
        "company": company,
        "job_title": title,
        "company_type": _text(row.get("企业性质")),
        "industry": _text(row.get("行业分类")),
        "city": _text(row.get("工作地点")),
        "cohort": _text(row.get("届次")),
        "education": _text(row.get("学历要求")),
        "updated_date": _text(row.get("更新时间")),
        "match_basis": _text(row.get("匹配依据")),
        "recommended_resume": _text(row.get("建议简历")),
        "risk_note": _text(row.get("风险提醒")),
        "announcement_url": _text(row.get("公告链接")),
        "application_url": _text(row.get("投递链接")),
        "source_status": _text(row.get("投递状态")),
        "personal_note": _text(row.get("个人备注")),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="导入校招总库与个人推进岗位池")
    parser.add_argument("--master", type=Path, required=True)
    parser.add_argument("--progress-pool", type=Path, required=True)
    args = parser.parse_args()

    all_entries = [
        _entry("all", "📁26、27校招汇总表", row_number, row)
        for row_number, row in _records(args.master, "📁26、27校招汇总表", 1)
        if _text(row.get("公司名称")) and _text(row.get("招聘岗位"))
    ]

    progress_by_identity: dict[str, dict] = {}
    for sheet_name, core in (("优先投递", True), ("全部可投岗位", False)):
        for row_number, row in _records(args.progress_pool, sheet_name, 4):
            company = _text(row.get("公司名称"))
            title = _text(row.get("推荐投递岗位"))
            if not company or not title:
                continue
            identity = _identity(company, title)
            if identity in progress_by_identity:
                continue
            progress_by_identity[identity] = _entry("progress", sheet_name, row_number, row, core=core)
    progress_entries = list(progress_by_identity.values())

    init_db()
    db = SessionLocal()
    try:
        db.query(JobLibraryEntry).filter(JobLibraryEntry.collection.in_(("all", "progress"))).delete(synchronize_session=False)
        for start in range(0, len(all_entries), 1000):
            db.bulk_insert_mappings(JobLibraryEntry, all_entries[start:start + 1000])
        db.bulk_insert_mappings(JobLibraryEntry, progress_entries)
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()

    print(json.dumps({
        "all": len(all_entries),
        "progress": len(progress_entries),
        "core_recommended": sum(item["recommendation_level"] == "核心推荐" for item in progress_entries),
        "applicable_only": sum(item["recommendation_level"] == "可投递" for item in progress_entries),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
