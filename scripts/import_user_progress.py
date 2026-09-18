from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
API_ROOT = PROJECT_ROOT / "apps" / "api"
if str(API_ROOT) not in sys.path:
    sys.path.insert(0, str(API_ROOT))

from app.core.database import SessionLocal, init_db  # noqa: E402
from app.services.application_import import import_application_progress  # noqa: E402


def _text(value) -> str:
    return str(value or "").strip()


def _applied_recommendations(path: Path) -> bytes:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["优先投递"]
        headers = [_text(cell.value) for cell in sheet[4]]
        rows: list[dict[str, object]] = []
        for values in sheet.iter_rows(min_row=5, values_only=True):
            source = {headers[index]: value for index, value in enumerate(values) if index < len(headers)}
            if _text(source.get("投递状态")) != "已投递":
                continue
            notes = "；".join(
                item for item in (
                    _text(source.get("匹配依据")),
                    _text(source.get("风险提醒")),
                    _text(source.get("个人备注")),
                ) if item
            )
            rows.append(
                {
                    "company": _text(source.get("公司名称")),
                    "job_title": _text(source.get("推荐投递岗位")),
                    "job_url": _text(source.get("投递链接") or source.get("公告链接")),
                    "city": _text(source.get("工作地点")),
                    "status": "applied",
                    "priority": 1,
                    "job_category": _text(source.get("目标方向") or source.get("主要方向")),
                    "application_channel": "推荐岗位清单",
                    "application_notes": notes,
                    "application_tags": ["核心推荐", "推荐清单已投递"],
                }
            )
        return json.dumps(rows, ensure_ascii=False).encode("utf-8")
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="合并用户的校招进度与推荐清单已投递记录")
    parser.add_argument("--database", type=Path, required=True)
    parser.add_argument("--progress", type=Path, required=True)
    parser.add_argument("--recommendations", type=Path, required=True)
    args = parser.parse_args()

    database = args.database.resolve()
    backup_dir = database.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = backup_dir / f"{database.stem}-before-progress-import-{stamp}{database.suffix}"
    shutil.copy2(database, backup)

    init_db()
    db = SessionLocal()
    try:
        manual_result = import_application_progress(
            db,
            filename=args.progress.name,
            content=args.progress.read_bytes(),
        )
        recommendation_result = import_application_progress(
            db,
            filename="推荐岗位清单-已投递.json",
            content=_applied_recommendations(args.recommendations),
        )
    finally:
        db.close()

    print(json.dumps({
        "backup": str(backup),
        "manual_progress": manual_result,
        "recommended_applied": recommendation_result,
    }, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
