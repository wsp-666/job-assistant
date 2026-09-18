import io
import unittest

from openpyxl import Workbook, load_workbook
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.database import Base, get_db
from app.models.models import ApplicationEvent, Job
from app.routers.applications import router as applications_router
from app.services.application_import import (
    ApplicationImportError,
    build_application_import_template,
    import_application_progress,
)
from app.services.job_pipeline import parse_tags


class ApplicationImportTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine)
        self.db = self.Session()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def test_csv_creates_then_updates_by_url_without_duplicates(self):
        first = (
            "公司,岗位,岗位链接,求职进度,优先级,标签,投递时间\n"
            "示例科技,AI FDE,https://jobs.example.com/1?utm_source=feishu,面试中,P0,AI、交付,2026-08-20\n"
        ).encode("utf-8-sig")
        result = import_application_progress(self.db, filename="飞书求职进度.csv", content=first)

        self.assertEqual(result["created"], 1)
        self.assertEqual(result["failed"], 0)
        job = self.db.query(Job).one()
        self.assertEqual(job.status, "interview")
        self.assertEqual(job.priority, 1)
        self.assertEqual(job.applied_at.strftime("%Y-%m-%d"), "2026-08-20")
        self.assertEqual(parse_tags(job.application_tags_json), ["AI", "交付"])

        second = (
            "公司,岗位,岗位链接,求职进度,优先级,备注\n"
            "示例科技,AI FDE,https://jobs.example.com/1,已获Offer,P1,薪资沟通中\n"
        ).encode("utf-8")
        result = import_application_progress(self.db, filename="更新.csv", content=second)

        self.assertEqual(result["updated"], 1)
        self.assertEqual(self.db.query(Job).count(), 1)
        self.db.refresh(job)
        self.assertEqual(job.status, "offer")
        self.assertEqual(job.priority, 2)
        self.assertEqual(job.application_notes, "薪资沟通中")
        self.assertGreaterEqual(
            self.db.query(ApplicationEvent).filter_by(job_id=job.id, event_type="progress_imported").count(),
            2,
        )

    def test_xlsx_import_warns_on_unknown_status_but_preserves_row(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["公司名称", "职位名称", "投递进度", "匹配度"])
        sheet.append(["云智公司", "解决方案工程师", "等待业务回复", "88%"])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()

        result = import_application_progress(self.db, filename="进度.xlsx", content=output.getvalue())

        self.assertEqual(result["created"], 1)
        self.assertEqual(result["failed"], 0)
        self.assertEqual(len(result["issues"]), 1)
        job = self.db.query(Job).one()
        self.assertEqual(job.status, "pending")
        self.assertEqual(job.match_score, 88)

    def test_shared_portal_url_does_not_merge_different_jobs(self):
        content = (
            "公司,岗位,投递链接,当前状态\n"
            "同一公司,Java开发,https://jobs.example.com/campus,已投递\n"
            "同一公司,产品经理,https://jobs.example.com/campus,已投递\n"
        ).encode("utf-8")
        result = import_application_progress(self.db, filename="共享门户.csv", content=content)

        self.assertEqual(result["created"], 2)
        self.assertEqual(self.db.query(Job).count(), 2)
        self.assertEqual({row.job_title for row in self.db.query(Job).all()}, {"Java开发", "产品经理"})

    def test_invalid_new_row_is_reported_without_aborting_valid_rows(self):
        content = (
            "公司,岗位,求职进度\n"
            ",,已投递\n"
            "有效公司,售前工程师,已投递\n"
        ).encode("utf-8")
        result = import_application_progress(self.db, filename="混合.csv", content=content)

        self.assertEqual(result["created"], 1)
        self.assertEqual(result["failed"], 1)
        self.assertEqual(self.db.query(Job).count(), 1)

    def test_template_has_expected_headers_and_unsupported_file_is_rejected(self):
        workbook = load_workbook(io.BytesIO(build_application_import_template()), read_only=True)
        self.assertEqual(workbook["求职进度"]["A1"].value, "序号")
        self.assertEqual(workbook["求职进度"]["B1"].value, "公司")
        self.assertEqual(workbook["求职进度"]["I1"].value, "当前状态")
        self.assertEqual(workbook["求职进度"]["K1"].value, "笔试")
        self.assertEqual(workbook["求职进度"]["L1"].value, "测评")
        self.assertIn("填写说明", workbook.sheetnames)
        workbook.close()

        with self.assertRaises(ApplicationImportError):
            import_application_progress(self.db, filename="旧格式.xls", content=b"not excel")

    def test_school_recruitment_columns_import_round_progress(self):
        workbook = Workbook()
        dashboard = workbook.active
        dashboard.title = "数据看板"
        dashboard.append(["投递数据看板"])
        sheet = workbook.create_sheet("投递进度")
        sheet.append(["校招投递进度管理"])
        sheet.append(["说明"])
        sheet.append([
            "序号", "公司", "链接", "岗位", "岗位大类", "城市", "投递渠道", "投递日期",
            "当前状态", "下一节点时间", "笔试/测评", "初筛", "一面", "二面", "三面", "终面",
            "薪资/备注要点", "更新日",
        ])
        sheet.append([
            1, "多岗位公司", "https://jobs.example.com/java", "Java开发", "开发", "深圳", "官网",
            "2026-09-01", "面试中", "2026-09-08", "通过", "通过", "通过", "待定", "未开始",
            "未开始", "薪资待沟通", "2026-09-02",
        ])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()

        result = import_application_progress(self.db, filename="校招进度管理.xlsx", content=output.getvalue())

        self.assertEqual(result["created"], 1)
        job = self.db.query(Job).one()
        self.assertEqual(job.job_category, "开发")
        self.assertEqual(job.written_test_status, "通过")
        self.assertEqual(job.assessment_status, "通过")
        self.assertEqual(job.second_interview_status, "待定")
        self.assertEqual(job.application_notes, "薪资待沟通")

    def test_http_upload_and_template_endpoints(self):
        app = FastAPI()
        app.include_router(applications_router)

        def override_db():
            yield self.db

        app.dependency_overrides[get_db] = override_db
        client = TestClient(app)
        response = client.post(
            "/api/applications/import",
            files={
                "file": (
                    "进度.csv",
                    "公司,岗位,求职进度\n接口公司,解决方案工程师,已投递\n".encode("utf-8"),
                    "text/csv",
                )
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["created"], 1)

        template = client.get("/api/applications/import-template")
        self.assertEqual(template.status_code, 200)
        self.assertGreater(len(template.content), 1000)
        self.assertIn("application-progress-template.xlsx", template.headers["content-disposition"])


if __name__ == "__main__":
    unittest.main()
