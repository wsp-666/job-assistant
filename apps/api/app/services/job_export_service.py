"""将已采集岗位导出为 Excel，并从 JD 中解析实习时间与技能要点。"""

import re
from datetime import datetime
from io import BytesIO
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.models.models import Job

COMMON_SKILL_KEYWORDS = [
    "Python",
    "Java",
    "JavaScript",
    "TypeScript",
    "Go",
    "Golang",
    "C++",
    "C#",
    "React",
    "Vue",
    "Angular",
    "Node.js",
    "Spring",
    "Django",
    "Flask",
    "FastAPI",
    "SQL",
    "MySQL",
    "PostgreSQL",
    "MongoDB",
    "Redis",
    "Kafka",
    "Docker",
    "Kubernetes",
    "K8s",
    "Linux",
    "Git",
    "AWS",
    "机器学习",
    "深度学习",
    "NLP",
    "计算机视觉",
    "数据分析",
    "Excel",
    "PPT",
    "Figma",
    "Photoshop",
    "产品经理",
    "运营",
    "市场营销",
]

DATE_RANGE_RE = re.compile(
    r"(\d{4}[\s年./-]*\d{1,2}[\s月./-]*\d{0,2})\s*[-—至到~]\s*"
    r"(\d{4}[\s年./-]*\d{1,2}[\s月./-]*\d{0,2})"
)
ARRIVAL_RE = re.compile(r"(到岗|入职|报道)[时间日]*[：:]\s*([^\n，。；;]{2,24})")
DURATION_RE = re.compile(r"实习(?:时长)?[：:]?\s*(\d+)\s*个?月")
SKILL_LINE_RE = re.compile(r"(?:熟悉|精通|掌握|了解|具备)[^。\n]{0,80}")
REQ_SECTION_RE = re.compile(r"(任职要求|岗位要求|任职资格)[：:\s]*([\s\S]{0,900})")


def extract_internship_dates(jd: str, fallback_created: datetime | None) -> tuple[str, str]:
    """从 JD 解析实习起止时间；解析不到时用采集日期作为开始时间。"""
    if not jd:
        if fallback_created:
            return fallback_created.strftime("%Y-%m-%d"), ""
        return "", ""

    match = DATE_RANGE_RE.search(jd)
    if match:
        return match.group(1).strip(), match.group(2).strip()

    start = ""
    arrival = ARRIVAL_RE.search(jd)
    if arrival:
        start = arrival.group(2).strip()
    elif fallback_created:
        start = fallback_created.strftime("%Y-%m-%d")

    end = ""
    duration = DURATION_RE.search(jd)
    if duration:
        end = f"约{duration.group(1)}个月"

    return start, end


def extract_skills_from_jd(jd: str) -> str:
    """从 JD 提取应注意的技能与要求，供 Excel「应注意技能」列使用。"""
    if not jd:
        return ""

    found: set[str] = set()
    jd_lower = jd.lower()
    for keyword in COMMON_SKILL_KEYWORDS:
        if keyword.lower() in jd_lower:
            found.add(keyword)

    for match in SKILL_LINE_RE.finditer(jd):
        line = match.group(0).strip()
        if len(line) > 4:
            found.add(line[:100])

    section = REQ_SECTION_RE.search(jd)
    if section:
        for line in re.split(r"[\n；;。]", section.group(2)):
            line = line.strip().lstrip("0123456789.-、）) ")
            if 4 < len(line) < 120 and any(
                word in line for word in ("熟悉", "精通", "掌握", "经验", "能力", "优先")
            ):
                found.add(line)

    if not found:
        return ""

    items = sorted(found, key=len)[:15]
    return "；".join(items)


def build_jobs_excel(jobs: Sequence[Job]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "岗位采集"

    headers = ["序号", "公司", "岗位", "开始时间", "结束时间", "jd", "应注意技能"]
    ws.append(headers)
    header_font = Font(bold=True)
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, job in enumerate(jobs, 1):
        start_time, end_time = extract_internship_dates(job.jd_text, job.created_at)
        skills = extract_skills_from_jd(job.jd_text)
        ws.append(
            [
                index,
                job.company or "",
                job.job_title or "",
                start_time,
                end_time,
                job.jd_text or "",
                skills,
            ]
        )

    widths = [6, 18, 22, 14, 14, 52, 32]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width

    for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
